"""Bounded, thread-scoped documents. The database, never local disk, owns them."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import struct
import zipfile
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import PurePosixPath
from typing import Literal
from uuid import UUID, uuid4
from xml.etree.ElementTree import iterparse

from pydantic import BaseModel, ConfigDict, Field

from arete.dataio.db import db_connection
from arete.dataio.db import transaction as db_transaction

MIB = 1024 * 1024
CHUNK_BYTES = 3 * MIB
MAX_FILE_BYTES = 20 * MIB
MAX_THREAD_BYTES = 100 * MIB
MAX_TOTAL_BYTES = 1024 * MIB
MAX_THREAD_FILES = 20
MAX_EXTRACT_BYTES = 2 * MIB
MAX_THREAD_EXTRACT_BYTES = 8 * MIB
MAX_CELLS = 100_000
MAX_SHEETS = 20
MAX_PAGES = 100
MAX_ARCHIVE_BYTES = 64 * MIB
MAX_ARCHIVE_ENTRIES = 1000
EXTENSIONS = frozenset(
    {".xlsx", ".xls", ".csv", ".md", ".txt", ".pdf", ".png", ".jpg", ".jpeg", ".webp"}
)


class DocumentError(ValueError):
    """An operating error that must remain visible to the athlete."""


@contextmanager
def document_transaction():
    """Serialize document lifecycle changes across processes, including deletion."""
    with db_transaction() as con:
        con.execute(
            "UPDATE app.document_quota SET used_bytes=used_bytes WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (id=1) "
        )
        yield con


class SourceBlock(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)
    locator: str = Field(min_length=1, max_length=200)
    text: str = Field(max_length=100_000)
    method: Literal["text", "ocr", "cell"] = "text"
    confidence: float | None = Field(default=None, ge=0, le=100)
    box: list[float] | None = Field(default=None, min_length=4, max_length=4)


class Extraction(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blocks: list[SourceBlock] = Field(min_length=1, max_length=MAX_CELLS)
    warnings: list[str] = Field(default_factory=list, max_length=100)


def _uuid(value: str) -> str:
    return str(UUID(value))


def _row(con, thread_id: str, document_id: str):
    row = con.execute(
        "SELECT id, name, size, sha256, status, extraction FROM app.visible_coach_documents WHERE id=? AND thread_id=?",
        [_uuid(document_id), _uuid(thread_id)],
    ).fetchone()
    if row is None:
        raise DocumentError("Document introuvable dans ce fil.")
    return row


def _metadata(row) -> dict:
    return dict(zip(("id", "name", "size", "sha256", "status"), row[:5], strict=True))


def list_documents(thread_id: str) -> list[dict]:
    with db_connection() as con:
        rows = con.execute(
            "SELECT id, name, size, sha256, status FROM app.visible_coach_documents WHERE thread_id=? ORDER BY created_at, id LIMIT ?",
            [_uuid(thread_id), MAX_THREAD_FILES + 1],
        ).fetchall()
    assert len(rows) <= MAX_THREAD_FILES
    return [_metadata(row) for row in rows]


def begin_upload(thread_id: str, name: str, size: int, sha256: str) -> dict:
    thread_id = _uuid(thread_id)
    if (
        not name
        or len(name) > 200
        or "/" in name
        or "\\" in name
        or any(ord(c) < 32 for c in name)
    ):
        raise DocumentError("Nom de fichier invalide.")
    if PurePosixPath(name).suffix.lower() not in EXTENSIONS:
        raise DocumentError("Format non pris en charge.")
    if not 0 < size <= MAX_FILE_BYTES:
        raise DocumentError("Le fichier doit contenir entre 1 octet et 20 Mio.")
    if len(sha256) != 64 or any(c not in "0123456789abcdef" for c in sha256):
        raise DocumentError("Empreinte SHA-256 invalide.")
    document_id = str(uuid4())
    with document_transaction() as con:
        # Updating one quota row serializes reservations across server instances.
        reserved = con.execute(
            "UPDATE app.document_quota SET used_bytes=used_bytes+? WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (id=1 AND used_bytes+?<=?) RETURNING id",
            [size, size, MAX_TOTAL_BYTES],
        ).fetchone()
        if not reserved:
            raise DocumentError(
                "Stockage des documents plein (1 Gio). Supprime des fichiers."
            )
        totals = con.execute(
            "SELECT count(*), coalesce(sum(size),0) FROM app.visible_coach_documents WHERE thread_id=?",
            [thread_id],
        ).fetchone()
        assert totals is not None
        count, used = totals
        if count >= MAX_THREAD_FILES or used + size > MAX_THREAD_BYTES:
            raise DocumentError("Limite du fil atteinte : 20 documents ou 100 Mio.")
        con.execute(
            "INSERT INTO app.coach_documents (id,thread_id,name,size,sha256,status) VALUES (?,?,?,?,?,'uploading')",
            [document_id, thread_id, name, size, sha256],
        )
    return {
        "id": document_id,
        "name": name,
        "size": size,
        "sha256": sha256,
        "status": "uploading",
    }


def put_chunk(thread_id: str, document_id: str, position: int, content: bytes) -> None:
    with document_transaction() as con:
        row = _row(con, thread_id, document_id)
        size = row[2]
        if (
            row[4] != "uploading"
            or not 0 <= position < (size + CHUNK_BYTES - 1) // CHUNK_BYTES
        ):
            raise DocumentError("Bloc inattendu ou document déjà finalisé.")
        if len(content) != min(CHUNK_BYTES, size - position * CHUNK_BYTES):
            raise DocumentError("Taille du bloc incorrecte.")
        old = con.execute(
            "SELECT content FROM app.visible_coach_document_chunks WHERE document_id=? AND position=?",
            [document_id, position],
        ).fetchone()
        if old:
            if bytes(old[0]) != content:
                raise DocumentError("Ce bloc existe avec un contenu différent.")
            return
        con.execute(
            "INSERT INTO app.coach_document_chunks (document_id,position,content) VALUES (?,?,?)",
            [document_id, position, content],
        )


def original(thread_id: str, document_id: str) -> tuple[str, bytes]:
    with db_connection() as con:
        row = _row(con, thread_id, document_id)
        chunks = con.execute(
            "SELECT position,content FROM app.visible_coach_document_chunks WHERE document_id=? ORDER BY position LIMIT 8",
            [document_id],
        ).fetchall()
    expected = (row[2] + CHUNK_BYTES - 1) // CHUNK_BYTES
    if len(chunks) != expected or [p for p, _ in chunks] != list(range(expected)):
        raise DocumentError("Téléversement incomplet.")
    raw = b"".join(bytes(c) for _, c in chunks)
    if len(raw) != row[2] or hashlib.sha256(raw).hexdigest() != row[3]:
        raise DocumentError("L’intégrité du document n’a pas pu être vérifiée.")
    return row[1], raw


def _cell(value) -> str:
    return value.isoformat() if isinstance(value, date | datetime) else str(value)


def _xls_formulas(book, sheet) -> dict[tuple[int, int], str]:
    """xlrd exposes cached values publicly; retain BIFF formula tokens as well.

    Never evaluate a formula. Its decompiler cannot recover every shared/array
    expression: preserve the raw tokens and an explicit warning in that case.
    The offsets and record format follow xlrd 2.x's Sheet.read implementation.
    """
    from xlrd.biffh import XL_FORMULA_OPCODES
    from xlrd.formula import FMLA_TYPE_CELL, decompile_formula

    result: dict[tuple[int, int], str] = {}
    position = sheet._position
    if book.mem is None:
        raise DocumentError(
            "Ce vieux format XLS ne conserve pas les formules ; convertis-le en XLSX."
        )
    for _ in range(MAX_CELLS * 4):
        if position + 4 > len(book.mem):
            raise DocumentError("Flux XLS interrompu.")
        code, length = struct.unpack_from("<HH", book.mem, position)
        position += 4
        data = book.mem[position : position + length]
        position += length
        if len(data) != length:
            raise DocumentError("Enregistrement XLS incomplet.")
        if code == 0x000A:  # Worksheet EOF.
            return result
        if code not in XL_FORMULA_OPCODES:
            continue
        if book.biff_version < 50 or length < 22:
            raise DocumentError(
                "Formule XLS ancienne non prise en charge ; convertis le classeur en XLSX."
            )
        row, col = struct.unpack_from("<HH", data)
        token_length = struct.unpack_from("<H", data, 20)[0]
        tokens = data[22 : 22 + token_length]
        if len(tokens) != token_length:
            raise DocumentError("Formule XLS incomplète.")
        try:
            expression = decompile_formula(
                book, tokens, token_length, FMLA_TYPE_CELL, browx=row, bcolx=col
            )
        except (ValueError, IndexError, KeyError, NotImplementedError):
            expression = None
        result[row, col] = (
            "=" + expression
            if expression and "SHARED" not in expression
            else f"FORMULE BIFF NON DÉCODÉE — à vérifier (jetons {tokens.hex()})"
        )
        if len(result) > MAX_CELLS:
            raise DocumentError("Maximum 100 000 cellules par classeur.")
    raise DocumentError("Trop d’enregistrements dans la feuille XLS.")


def extract_text(name: str, raw: bytes) -> Extraction | None:
    suffix = PurePosixPath(name).suffix.lower()
    blocks: list[SourceBlock] = []
    warnings: list[str] = []
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > MAX_ARCHIVE_ENTRIES
                or sum(e.file_size for e in entries) > MAX_ARCHIVE_BYTES
            ):
                raise DocumentError("Classeur décompressé trop volumineux.")
            if any("vbaProject" in e.filename for e in entries):
                raise DocumentError("Les macros ne sont pas acceptées.")
            sheets = [
                e
                for e in entries
                if e.filename.startswith("xl/worksheets/sheet")
                and e.filename.endswith(".xml")
            ]
            if len(sheets) > MAX_SHEETS:
                raise DocumentError("Maximum 20 feuilles par classeur.")
            cells = 0
            # Count before openpyxl allocates cell objects, including sparse sheets.
            for entry in sheets:
                with archive.open(entry) as stream:
                    for _, element in iterparse(stream, events=("end",)):
                        if element.tag.endswith("}c"):
                            cells += 1
                            if cells > MAX_CELLS:
                                raise DocumentError(
                                    "Maximum 100 000 cellules par classeur."
                                )
                        element.clear()
        workbook = load_workbook(
            io.BytesIO(raw), read_only=False, data_only=False, keep_links=False
        )
        values = load_workbook(
            io.BytesIO(raw), read_only=True, data_only=True, keep_links=False
        )
        try:
            if len(workbook.worksheets) > MAX_SHEETS:
                raise DocumentError("Maximum 20 feuilles par classeur.")
            if sum(s.max_row * s.max_column for s in workbook.worksheets) > MAX_CELLS:
                raise DocumentError("Maximum 100 000 cellules par classeur.")
            for sheet in workbook.worksheets:
                cached_values = {
                    cell.coordinate: cell.value
                    for row in values[sheet.title].iter_rows()
                    for cell in row
                    if cell.value is not None
                }
                for merged in sheet.merged_cells.ranges:
                    warnings.append(f"Cellules fusionnées : {sheet.title}!{merged}")
                for cells in sheet.iter_rows():
                    for cell in cells:
                        if cell.value is None:
                            continue
                        text = _cell(cell.value)
                        if cell.data_type == "f":
                            cached = cached_values.get(cell.coordinate)
                            text += f" [valeur enregistrée : {_cell(cached) if cached is not None else 'MANQUANTE — à vérifier'}]"
                        blocks.append(
                            SourceBlock(
                                locator=f"{sheet.title}!{cell.coordinate}",
                                text=text,
                                method="cell",
                            )
                        )
        finally:
            workbook.close()
            values.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(
            file_contents=raw,
            on_demand=True,
            formatting_info=True,
            logfile=io.StringIO(),
        )
        try:
            if (
                book.nsheets > MAX_SHEETS
                or sum(s.nrows * s.ncols for s in book.sheets()) > MAX_CELLS
            ):
                raise DocumentError(
                    "Classeur trop volumineux (20 feuilles, 100 000 cellules)."
                )
            warnings.append(
                "XLS : formules et valeurs enregistrées, jamais recalculées. Les formules BIFF non décodées exigent une vérification dans l’original."
            )
            for sheet in book.sheets():
                formulas = _xls_formulas(book, sheet)
                for rlo, rhi, clo, chi in sheet.merged_cells:
                    warnings.append(
                        f"Cellules fusionnées : {sheet.name}!R{rlo + 1}C{clo + 1}:R{rhi}C{chi}"
                    )
                for r in range(sheet.nrows):
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        if cell.ctype == xlrd.XL_CELL_EMPTY:
                            continue
                        value = (
                            xlrd.xldate_as_datetime(cell.value, book.datemode)
                            if cell.ctype == xlrd.XL_CELL_DATE
                            else cell.value
                        )
                        text = _cell(value)
                        if (r, c) in formulas:
                            text = f"{formulas[r, c]} [valeur enregistrée : {text if text else 'MANQUANTE — à vérifier'}]"
                        blocks.append(
                            SourceBlock(
                                locator=f"{sheet.name}!R{r + 1}C{c + 1}",
                                text=text,
                                method="cell",
                            )
                        )
        finally:
            book.release_resources()
    elif suffix in {".txt", ".md", ".csv"}:
        text = raw.decode("utf-8-sig")
        if "\x00" in text:
            raise DocumentError("Le fichier contient des données binaires.")
        if suffix == ".csv":
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            count = 0
            for row_number, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
                count += len(row)
                if count > MAX_CELLS or row_number > MAX_CELLS:
                    raise DocumentError("Maximum 100 000 cellules par CSV.")
                blocks.extend(
                    SourceBlock(
                        locator=f"R{row_number}C{c + 1}", text=value, method="cell"
                    )
                    for c, value in enumerate(row)
                    if value
                )
        else:
            # Chunks have stable line references and do not alter the source text.
            lines = text.splitlines()
            for start in range(0, len(lines), 100):
                blocks.append(
                    SourceBlock(
                        locator=f"lignes {start + 1}-{min(start + 100, len(lines))}",
                        text="\n".join(lines[start : start + 100]),
                    )
                )
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise DocumentError("Les PDF chiffrés ne sont pas acceptés.")
        if len(reader.pages) > MAX_PAGES:
            raise DocumentError("Maximum 100 pages par PDF.")
        return None  # The browser provides position-aware text/OCR, never a model call.
    else:
        from PIL import Image

        with Image.open(io.BytesIO(raw)) as img:
            if (
                img.format
                != {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[
                    suffix
                ]
                or img.width * img.height > 25_000_000
            ):
                raise DocumentError("Image invalide ou trop grande (25 mégapixels).")
            img.verify()
        return None
    return Extraction(blocks=blocks, warnings=warnings)


def finalize(
    thread_id: str, document_id: str, browser_extraction: Extraction | None = None
) -> dict:
    name, raw = original(thread_id, document_id)
    try:
        extraction = extract_text(name, raw)
    except Exception as exc:
        raise DocumentError(f"Lecture du document impossible : {exc}") from exc
    if extraction is None:
        if browser_extraction is None:
            raise DocumentError("L’extraction du PDF ou de l’image est manquante.")
        extraction = browser_extraction
    encoded = extraction.model_dump_json()
    if len(encoded.encode()) > MAX_EXTRACT_BYTES:
        raise DocumentError("Extraction trop volumineuse (2 Mio). Sépare le document.")
    if not any(b.text.strip() for b in extraction.blocks):
        raise DocumentError("Aucun texte lisible dans ce document.")
    with document_transaction() as con:
        row = _row(con, thread_id, document_id)
        if row[4] == "ready":
            return _metadata(row)
        totals = con.execute(
            "SELECT coalesce(sum(octet_length(encode(CAST(extraction AS VARCHAR)))),0) FROM app.visible_coach_documents WHERE thread_id=?",
            [thread_id],
        ).fetchone()
        assert totals is not None
        used = totals[0]
        if used + len(encoded.encode()) > MAX_THREAD_EXTRACT_BYTES:
            raise DocumentError("Les extractions de ce fil dépassent 8 Mio.")
        con.execute(
            "UPDATE app.coach_documents SET status='ready', extraction=? WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (id=?) ",
            [encoded, document_id],
        )
    return {**_metadata(row), "status": "ready"}


def extraction_for(thread_id: str, document_id: str) -> Extraction:
    with db_connection() as con:
        row = _row(con, thread_id, document_id)
    if row[4] != "ready":
        raise DocumentError("Le document n’est pas prêt.")
    return Extraction.model_validate_json(row[5])


def filesystem(
    thread_id: str, document_ids: tuple[str, ...] | None = None
) -> dict[str, str]:
    with db_connection() as con:
        rows = con.execute(
            "SELECT id,name,extraction,status FROM app.visible_coach_documents WHERE thread_id=? ORDER BY created_at, id LIMIT ?",
            [_uuid(thread_id), MAX_THREAD_FILES + 1],
        ).fetchall()
    assert len(rows) <= MAX_THREAD_FILES
    if document_ids is not None:
        selected = set(document_ids)
        available = {row[0] for row in rows}
        if len(selected) > MAX_THREAD_FILES or not selected.issubset(available):
            raise DocumentError(
                "Une pièce jointe sélectionnée est absente ou incomplète. Vérifie les fichiers du fil."
            )
        rows = [row for row in rows if row[0] in selected]
    if any(row[3] != "ready" for row in rows):
        # Legacy callers select the whole thread. Never quietly drop a file
        # whose parsing/OCR has not completed and answer without its evidence.
        raise DocumentError(
            "Une pièce jointe est incomplète. Termine son extraction ou retire-la avant de répondre."
        )
    files = {}
    total = 0
    for document_id, name, raw, _ in rows:
        extraction = Extraction.model_validate_json(raw)
        total += len(raw.encode())
        if total > MAX_THREAD_EXTRACT_BYTES:
            raise DocumentError("Les documents du fil dépassent le budget filesystem.")
        files[f"/attachments/{document_id}.md"] = (
            f"# {name}\nDocument : {document_id}\n"
            + "\n".join(extraction.warnings)
            + "\n\n"
            + "\n\n".join(
                f"[{b.locator}] ({b.method})\n{b.text}" for b in extraction.blocks
            )
        )
    return files


def delete_documents(thread_id: str, document_id: str | None = None) -> None:
    with document_transaction() as con:
        con.execute(
            "UPDATE app.document_quota SET used_bytes=used_bytes WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (id=1) "
        )
        where = "thread_id=?" + (" AND id=?" if document_id else "")
        params = [_uuid(thread_id)] + ([_uuid(document_id)] if document_id else [])
        rows = con.execute(
            f"SELECT id,size FROM app.visible_coach_documents WHERE {where}", params
        ).fetchall()
        for identifier, _ in rows:
            con.execute(
                "DELETE FROM app.coach_document_chunks WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (document_id=?) ",
                [identifier],
            )
        con.execute(
            f"DELETE FROM app.coach_documents WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND ({where}) ",
            params,
        )
        con.execute(
            "UPDATE app.document_quota SET used_bytes=used_bytes-? WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (id=1) ",
            [sum(r[1] for r in rows)],
        )
        if document_id is None:
            con.execute(
                "DELETE FROM app.coach_imports WHERE athlete_id = getvariable('arete_athlete_id') AND deleted_at IS NULL AND EXISTS (SELECT 1 FROM app.athletes scope_owner WHERE scope_owner.id=athlete_id AND scope_owner.deleted_at IS NULL) AND (thread_id=?) ",
                [thread_id],
            )


def manifest(thread_id: str, document_ids: tuple[str, ...] | None = None) -> str:
    return json.dumps(
        [
            d
            for d in list_documents(thread_id)
            if document_ids is None or d["id"] in document_ids
        ],
        ensure_ascii=False,
    )

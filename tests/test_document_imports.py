"""Document/import boundaries: durable bytes, provenance and human-only writes."""

import asyncio
import hashlib
import io
import json
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from uuid import uuid4

import duckdb
import pytest
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from arete.dataio.db import db_connection
from arete.dataio.init_duckdb import main
from arete.services import documents, imports
from arete.services.documents import DocumentError, Extraction, SourceBlock
from arete.services.prescriptions import ImportedSession, Prescription, Provenance, Step


@pytest.fixture
def document_db(tmp_path, monkeypatch):
    monkeypatch.setenv("ARETE_DB", str(tmp_path / "documents.duckdb"))
    main()
    return str(uuid4())


def uploaded(
    thread, name="plan.md", raw=b"2027-01-12 : footing 30 minutes", extraction=None
):
    doc = documents.begin_upload(
        thread, name, len(raw), hashlib.sha256(raw).hexdigest()
    )
    for position in range(
        (len(raw) + documents.CHUNK_BYTES - 1) // documents.CHUNK_BYTES
    ):
        documents.put_chunk(
            thread,
            doc["id"],
            position,
            raw[
                position * documents.CHUNK_BYTES : (position + 1)
                * documents.CHUNK_BYTES
            ],
        )
    return documents.finalize(thread, doc["id"], extraction)


def session_for(doc):
    return ImportedSession(
        date=date(2027, 1, 12),
        sport="running",
        description="Footing",
        prescription=Prescription(
            steps=[Step(kind="effort", duration_kind="seconds", value=1800)]
        ),
        provenance=[
            Provenance(
                document_id=doc["id"], locator="lignes 1-1", quote="footing 30 minutes"
            )
        ],
    )


def test_upload_integrity_retry_and_thread_isolation(document_db):
    thread = document_db
    raw = b"plan"
    doc = documents.begin_upload(thread, "plan.md", 4, hashlib.sha256(raw).hexdigest())
    with pytest.raises(DocumentError, match="incomplet"):
        documents.finalize(thread, doc["id"])
    documents.put_chunk(thread, doc["id"], 0, raw)
    documents.put_chunk(thread, doc["id"], 0, raw)
    with pytest.raises(DocumentError, match="différent"):
        documents.put_chunk(thread, doc["id"], 0, b"evil")
    with pytest.raises(DocumentError, match="introuvable"):
        documents.original(str(uuid4()), doc["id"])
    documents.finalize(thread, doc["id"])
    main()  # A new boot does not discard files.
    assert documents.original(thread, doc["id"])[1] == raw
    assert "plan" in documents.filesystem(thread)[f"/attachments/{doc['id']}.md"]


def test_bad_digest_never_marks_document_ready(document_db):
    doc = documents.begin_upload(document_db, "p.txt", 4, "0" * 64)
    documents.put_chunk(document_db, doc["id"], 0, b"plan")
    with pytest.raises(DocumentError, match="intégrité"):
        documents.finalize(document_db, doc["id"])
    assert documents.list_documents(document_db)[0]["status"] == "uploading"


def test_quota_reserves_before_upload_and_releases_on_delete(document_db, monkeypatch):
    monkeypatch.setattr(documents, "MAX_TOTAL_BYTES", 5)
    first = documents.begin_upload(document_db, "p.md", 4, "0" * 64)
    with pytest.raises(DocumentError, match="plein"):
        documents.begin_upload(document_db, "q.md", 2, "0" * 64)
    documents.delete_documents(document_db, first["id"])
    documents.begin_upload(document_db, "q.md", 5, "0" * 64)
    with pytest.raises(DocumentError, match="Nom"):
        documents.begin_upload(document_db, "../tokens.md", 1, "0" * 64)


def test_real_workbook_preserves_dates_formulas_and_merged_cells(document_db):
    book = Workbook()
    sheet = book.active
    sheet.title = "Semaine 1"
    sheet.append([date(2027, 1, 12), "Course", 30])
    sheet["D1"] = "=C1*60"
    sheet.merge_cells("A2:B2")
    sheet["A2"] = "Repos"
    book.create_sheet("Musculation").append(["Squat", "3 x 8", "60 kg"])
    output = io.BytesIO()
    book.save(output)
    doc = uploaded(document_db, "plan.xlsx", output.getvalue())
    extraction = documents.extraction_for(document_db, doc["id"])
    assert any(
        b.locator == "Semaine 1!A1" and "2027-01-12" in b.text
        for b in extraction.blocks
    )
    assert any(
        b.locator == "Semaine 1!D1" and "MANQUANTE" in b.text for b in extraction.blocks
    )
    assert any("fusionnées" in warning for warning in extraction.warnings)
    assert any(
        b.locator == "Musculation!C1" and b.text == "60 kg" for b in extraction.blocks
    )


def test_csv_french_units_and_no_silent_cell_truncation(document_db, monkeypatch):
    raw = b"date;distance;allure\n12/01/2027;8,5 km;5:30/km"
    doc = uploaded(document_db, "plan.csv", raw)
    assert any(
        b.text == "8,5 km"
        for b in documents.extraction_for(document_db, doc["id"]).blocks
    )
    monkeypatch.setattr(documents, "MAX_CELLS", 3)
    with pytest.raises(DocumentError, match="cellules"):
        documents.extract_text("plan.csv", raw)


def test_scanned_image_preserves_ocr_provenance_and_pdf_rejects_encryption(document_db):
    image = Image.new("RGB", (20, 20), "white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    doc = uploaded(
        document_db,
        "scan.png",
        buffer.getvalue(),
        Extraction(
            blocks=[
                SourceBlock(
                    locator="page 1",
                    text="Course 30 minutes",
                    method="ocr",
                    confidence=51,
                    box=[0, 0, 20, 20],
                )
            ]
        ),
    )
    assert documents.extraction_for(document_db, doc["id"]).blocks[0].confidence == 51
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentError, match="chiffrés"):
        uploaded(document_db, "locked.pdf", buffer.getvalue())


def test_preview_provenance_confirmation_and_deletion_preserve_sessions(document_db):
    doc = uploaded(document_db)
    proposal = session_for(doc)
    assert imports.has_unvalidated_documents(document_db)
    preview = imports.propose(document_db, [proposal])
    with db_connection() as con:
        assert (
            con.execute("SELECT count(*) FROM app.planned_sessions").fetchone()[0] == 0
        )
    with pytest.raises(DocumentError, match="Vérifie"):
        imports.confirm(document_db, preview["draft_id"], 1, str(uuid4()), [0], False)
    key = str(uuid4())
    ids = imports.confirm(document_db, preview["draft_id"], 1, key, [0], True)
    assert imports.confirm(document_db, preview["draft_id"], 1, key, [0], True) == ids
    assert not imports.has_unvalidated_documents(document_db)
    documents.delete_documents(document_db)
    with db_connection() as con:
        row = con.execute(
            "SELECT prescription,provenance FROM app.planned_sessions WHERE id=?", ids
        ).fetchone()
    assert json.loads(row[0])["steps"][0]["value"] == 1800
    assert json.loads(row[1])[0]["quote"] == "footing 30 minutes"


def test_false_provenance_and_ambiguous_dates_cannot_commit(document_db):
    doc = uploaded(document_db)
    proposal = session_for(doc)
    proposal.provenance[0].quote = "invented"
    with pytest.raises(DocumentError, match="extrait"):
        imports.propose(document_db, [proposal])
    proposal = session_for(doc).model_copy(update={"date": None})
    draft = imports.propose(document_db, [proposal])
    with pytest.raises(DocumentError, match="problèmes"):
        imports.confirm(document_db, draft["draft_id"], 1, str(uuid4()), [0], True)
    imports.update_draft(document_db, draft["draft_id"], 1, [session_for(doc)])
    with pytest.raises(DocumentError, match="version"):
        imports.confirm(document_db, draft["draft_id"], 1, str(uuid4()), [0], True)


def test_concurrent_confirmation_has_one_commit(document_db):
    doc = uploaded(document_db)
    preview = imports.propose(document_db, [session_for(doc)])
    key = str(uuid4())

    def run():
        try:
            return imports.confirm(document_db, preview["draft_id"], 1, key, [0], True)
        except (DocumentError, duckdb.TransactionException):
            return None

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: run(), range(2)))
    assert any(results)
    with db_connection() as con:
        assert (
            con.execute("SELECT count(*) FROM app.planned_sessions").fetchone()[0] == 1
        )


def test_statebackend_reads_documents_and_rejects_writes(
    document_db, monkeypatch, tmp_path
):
    from langchain.agents import create_agent
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
    from langchain_core.messages import AIMessage, HumanMessage

    from arete.agent.backends.memory import build_memory_filesystem
    from arete.agent.runtime.context import AgentContext
    from arete.api.agent import _document_state

    monkeypatch.setattr("arete.agent.backends.memory.memory_root", lambda: tmp_path)
    doc = uploaded(document_db)

    class Model(GenericFakeChatModel):
        def bind_tools(self, tools, **kwargs):
            return self

    messages = [
        AIMessage(
            content="",
            tool_calls=[
                {
                    "name": "read_file",
                    "args": {"file_path": f"/attachments/{doc['id']}.md"},
                    "id": "read",
                }
            ],
        ),
        AIMessage(content="OK"),
    ]
    context = AgentContext(thread_id=document_db)
    state, _ = _document_state(context)
    filesystem = build_memory_filesystem()
    assert {t.name for t in filesystem.tools} == {
        "read_file",
        "ls",
        "glob",
        "grep",
        "edit_file",
        "delete",
    }
    for tool in filesystem.tools:
        assert "anyOf" not in json.dumps(tool.args_schema.model_json_schema())
    graph = create_agent(
        Model(messages=iter(messages)),
        middleware=[filesystem],
        context_schema=AgentContext,
    )
    result = graph.invoke(
        {"messages": [HumanMessage("Lis le fichier")], **state}, context=context
    )
    assert "footing 30 minutes" in next(
        m.content for m in result["messages"] if m.type == "tool"
    )
    other, _ = _document_state(AgentContext(thread_id=str(uuid4())))
    assert other["files"] == {}


@pytest.mark.parametrize("async_mode", [False, True])
def test_attachment_grep_returns_matching_lines_by_default(
    tmp_path, monkeypatch, async_mode
):
    from deepagents.backends.utils import create_file_data
    from langchain.agents import create_agent
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
    from langchain_core.messages import AIMessage, HumanMessage

    from arete.agent.backends.memory import build_memory_filesystem
    from arete.agent.runtime.context import AgentContext

    monkeypatch.setattr("arete.agent.backends.memory.memory_root", lambda: tmp_path)

    class Model(GenericFakeChatModel):
        def bind_tools(self, tools, **kwargs):
            return self

    path = "/attachments/plan.md"
    args = {"path": path, "pattern": "2026-10-12", "max_count": 1}
    graph = create_agent(
        Model(
            messages=iter(
                [
                    AIMessage(
                        content="",
                        tool_calls=[
                            {"name": "grep", "args": args, "id": "content"},
                            {
                                "name": "grep",
                                "args": {**args, "output_mode": "files_with_matches"},
                                "id": "files",
                            },
                        ],
                    ),
                    AIMessage(content="Fin."),
                ]
            )
        ),
        middleware=[build_memory_filesystem()],
        context_schema=AgentContext,
    )
    state = {
        "messages": [HumanMessage("Lis la semaine du 12 octobre")],
        "files": {
            "/plan.md": create_file_data(
                "[Feuil1!B40]\n2026-10-12\n[Feuil1!A43]\n2026-10-12"
            )
        },
    }
    result = (
        asyncio.run(graph.ainvoke(state, context=AgentContext()))
        if async_mode
        else graph.invoke(state, context=AgentContext())
    )
    outputs = {m.tool_call_id: m for m in result["messages"] if m.type == "tool"}
    assert all(m.status == "success" for m in outputs.values())
    assert path in outputs["content"].content
    assert "2: 2026-10-12" in outputs["content"].content
    assert "4: 2026-10-12" not in outputs["content"].content
    assert "incomplete" in outputs["content"].content
    assert path in outputs["files"].content
    assert "2026-10-12" not in outputs["files"].content


def test_existing_planning_writes_blocked_during_import(document_db):
    from types import SimpleNamespace

    from arete.agent.capabilities.execution import _resolve_tool
    from arete.agent.runtime.context import AgentContext

    request = SimpleNamespace(
        tool_call={"name": "create_planned_session", "id": "write"},
        state={},
        runtime=SimpleNamespace(
            context=AgentContext(thread_id=document_db, document_import_pending=True)
        ),
    )
    assert _resolve_tool(request).status == "error"


def test_duplicate_batch_is_visible_and_discard_preserves_documents(document_db):
    doc = uploaded(document_db)
    preview = imports.propose(document_db, [session_for(doc), session_for(doc)])
    draft = imports.list_drafts(document_db)[0]
    assert draft["sessions"][0]["batch_duplicates"] == [2]
    assert draft["sessions"][1]["batch_duplicates"] == [1]
    imports.discard(document_db, preview["draft_id"], preview["version"])
    assert documents.list_documents(document_db)
    assert not imports.has_unvalidated_documents(document_db)
    with pytest.raises(DocumentError, match="version"):
        imports.confirm(
            document_db,
            preview["draft_id"],
            preview["version"],
            str(uuid4()),
            [0],
            True,
        )


def test_http_chunk_upload_and_thread_deletion(document_db):
    from fastapi.testclient import TestClient

    from arete.api.main import app

    root = f"/agent/threads/{document_db}"
    raw = b"2027-01-12 Course 30 minutes"
    with TestClient(app) as client:
        response = client.post(
            f"{root}/documents",
            json={
                "name": "plan.md",
                "size": len(raw),
                "sha256": hashlib.sha256(raw).hexdigest(),
            },
        )
        assert response.status_code == 201, response.text
        doc = response.json()
        path = f"{root}/documents/{doc['id']}"
        assert client.put(f"{path}/chunks/0", content=raw).status_code == 200
        assert client.post(f"{path}/finalize").status_code == 200
        assert client.get(f"{path}/chunks/0").content == raw
        assert (
            client.get(f"{path}/extraction").json()["blocks"][0]["text"] == raw.decode()
        )
        assert client.delete(root).status_code == 200
        assert client.get(f"{root}/documents").json() == []


def test_legacy_xls_retains_formula_and_missing_cache():
    from pathlib import Path

    extraction = documents.extract_text(
        "formulas.xls", Path("tests/data/imports/formulas.xls").read_bytes()
    )
    formula = next(b for b in extraction.blocks if b.locator == "Séances!R1C3")
    assert "B1*60" in formula.text
    assert "MANQUANTE" in formula.text
    assert any("fusionnées" in w for w in extraction.warnings)


def test_import_tool_uses_server_thread_and_never_confirms(document_db):
    from arete.agent.tools.imports import prepare_import

    doc = uploaded(document_db)
    session = session_for(doc)
    assert "config" not in prepare_import.args
    result = json.loads(
        prepare_import.invoke(
            {"sessions_json": json.dumps([session.model_dump(mode="json")])},
            config={"configurable": {"thread_id": document_db}},
        )
    )
    assert result["draft_id"] == imports.list_drafts(document_db)[0]["id"]
    with db_connection() as con:
        assert con.execute("SELECT count(*) FROM app.planned_sessions").fetchone() == (
            0,
        )
    other = str(uuid4())
    rejected = json.loads(
        prepare_import.invoke(
            {"sessions_json": json.dumps([session.model_dump(mode="json")])},
            config={"configurable": {"thread_id": other}},
        )
    )
    assert "error" in rejected
    assert imports.list_drafts(other) == []


def test_selected_documents_are_thread_scoped_and_never_silently_dropped(document_db):
    first = uploaded(document_db, "first.md", b"first")
    second = uploaded(document_db, "second.md", b"second")
    files = documents.filesystem(document_db, (second["id"],))
    assert list(files) == [f"/attachments/{second['id']}.md"]
    assert documents.filesystem(document_db, ()) == {}
    assert first["id"] not in documents.manifest(document_db, (second["id"],))
    with pytest.raises(DocumentError, match="absente"):
        documents.filesystem(str(uuid4()), (second["id"],))
    pending = documents.begin_upload(document_db, "pending.txt", 2, "0" * 64)
    with pytest.raises(DocumentError, match="incomplète"):
        documents.filesystem(document_db, (pending["id"],))
    with pytest.raises(DocumentError, match="incomplète"):
        documents.filesystem(document_db)
    # An unrelated unfinished upload must not block an explicit ready selection.
    assert documents.filesystem(document_db, (second["id"],)) == files


@pytest.mark.parametrize("streaming", [False, True])
@pytest.mark.parametrize("source", ["xlsx", "ocr"])
def test_chat_hydrates_extraction_before_first_model_and_reads_mounted_file(
    document_db, monkeypatch, tmp_path, router_client, streaming, source
):
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
    from langchain_core.messages import AIMessage

    from arete.agent.backends.memory import build_memory_filesystem
    from arete.agent.factory import build_agent
    from arete.agent.profiles.catalog import get_profile
    from arete.api.agent import router

    monkeypatch.setattr("arete.agent.backends.memory.memory_root", lambda: tmp_path)
    if source == "xlsx":
        book = Workbook()
        book.active.append([date(2026, 10, 12), "Footing après le 11", "45 minutes"])
        raw = io.BytesIO()
        book.save(raw)
        doc = uploaded(document_db, "Prépa Semi Lille.xlsx", raw.getvalue())
        locator = "A1"
    else:
        raw = io.BytesIO()
        Image.new("RGB", (10, 10)).save(raw, format="PNG")
        doc = uploaded(
            document_db,
            "scan.png",
            raw.getvalue(),
            Extraction(
                blocks=[
                    SourceBlock(
                        locator="page 1",
                        text="2026-10-12 : Footing après le 11, 45 minutes",
                        method="ocr",
                        confidence=90,
                    )
                ]
            ),
        )
        locator = "page 1"
    path = f"/attachments/{doc['id']}.md"
    seen = []

    class Model(GenericFakeChatModel):
        def bind_tools(self, tools, **kwargs):
            return self

        def _generate(self, messages, stop=None, run_manager=None, **kwargs):
            seen.append(messages)
            return super()._generate(
                messages, stop=stop, run_manager=run_manager, **kwargs
            )

    graph = build_agent(
        get_profile("chat"),
        model=Model(
            messages=iter(
                [
                    AIMessage(
                        content="",
                        tool_calls=[
                            {
                                "name": "read_file",
                                "args": {"file_path": path},
                                "id": "source",
                            }
                        ],
                    ),
                    AIMessage(content="Le fichier contient une séance le 12 octobre."),
                    AIMessage(content="Autre fil."),
                ]
            ),
            disable_streaming=True,
        ),
        context_tokens=65_536,
        output_tokens=4096,
        filesystem=build_memory_filesystem(),
    )
    monkeypatch.setattr("arete.api.agent.get_agent", lambda: graph)
    client = router_client(router)
    endpoint = "/agent/chat/stream" if streaming else "/agent/chat"
    response = client.post(
        endpoint,
        json={
            "thread_id": document_db,
            "document_ids": [doc["id"]],
            "messages": [
                {"role": "user", "content": "Exporte les séances après le 11 octobre"}
            ],
        },
    )
    assert response.status_code == 200
    assert "Le fichier contient" in response.text
    # Even a model that answers without reading tools receives actual source
    # evidence on its first call, including locators and OCR/cell provenance.
    assert path in seen[0][0].text
    assert "Footing après le 11" in seen[0][0].text
    assert locator in seen[0][0].text
    tool_result = next(m for m in seen[1] if m.type == "tool")
    assert tool_result.status == "success"
    assert "45 minutes" in tool_result.text
    assert len(seen) == 2
    client.post(
        endpoint,
        json={
            "thread_id": str(uuid4()),
            "document_ids": [],
            "messages": [{"role": "user", "content": "Autre conversation"}],
        },
    )
    assert path not in seen[2][0].text
    assert "Footing après le 11" not in seen[2][0].text


@pytest.mark.parametrize("streaming", [False, True])
@pytest.mark.parametrize(
    "failure", ["pending", "legacy_pending", "missing", "no_thread"]
)
def test_chat_rejects_unavailable_files_before_agent_creation(
    document_db, monkeypatch, router_client, streaming, failure
):
    from arete.api.agent import router

    doc = documents.begin_upload(document_db, "plan.xlsx", 4, "0" * 64)
    body = {
        "thread_id": document_db,
        "document_ids": [str(uuid4()) if failure == "missing" else doc["id"]],
        "messages": [{"role": "user", "content": "Lis mon fichier"}],
    }
    if failure == "legacy_pending":
        del body["document_ids"]
    if failure == "no_thread":
        del body["thread_id"]

    def unexpected_graph():
        pytest.fail("Agent must not start before document hydration succeeds")

    monkeypatch.setattr("arete.api.agent.get_agent", unexpected_graph)
    response = router_client(router).post(
        "/agent/chat/stream" if streaming else "/agent/chat", json=body
    )
    if streaming:
        events = [
            json.loads(line[6:])
            for line in response.text.splitlines()
            if line.startswith("data: ")
        ]
        assert [e["type"] for e in events] == ["error"]
    else:
        assert response.status_code == 422


def test_concurrent_threads_keep_their_own_attachment_context(
    document_db, monkeypatch, tmp_path
):
    from langchain.agents import create_agent
    from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
    from langchain_core.messages import AIMessage, HumanMessage

    from arete.agent.backends.memory import build_memory_filesystem
    from arete.agent.middlewares.context import ContextBuilderMiddleware
    from arete.agent.runtime.context import AgentContext
    from arete.api.agent import _document_state

    monkeypatch.setattr("arete.agent.backends.memory.memory_root", lambda: tmp_path)
    contexts = [
        AgentContext(thread_id=document_db),
        AgentContext(thread_id=str(uuid4())),
    ]
    markers = ["SOURCE_ALPHA", "SOURCE_BETA"]
    states = []
    for context, marker in zip(contexts, markers, strict=True):
        uploaded(context.thread_id, raw=marker.encode())
        state, _ = _document_state(context)
        states.append({"messages": [HumanMessage(marker)], **state})

    async def concurrent_runs():
        barrier = asyncio.Barrier(2)

        class Model(GenericFakeChatModel):
            def bind_tools(self, tools, **kwargs):
                return self

            async def _agenerate(self, messages, **kwargs):
                await asyncio.wait_for(barrier.wait(), timeout=5)
                own = messages[-1].text
                other = next(marker for marker in markers if marker != own)
                assert own in messages[0].text
                assert other not in messages[0].text
                return self._generate(messages)

        graph = create_agent(
            Model(messages=iter([AIMessage(content="OK"), AIMessage(content="OK")])),
            middleware=[build_memory_filesystem(), ContextBuilderMiddleware()],
            context_schema=AgentContext,
        )
        results = await asyncio.wait_for(
            asyncio.gather(
                *[
                    graph.ainvoke(state, context=context)
                    for state, context in zip(states, contexts, strict=True)
                ]
            ),
            timeout=10,
        )
        assert results[0]["files"] == states[0]["files"]
        assert results[1]["files"] == states[1]["files"]

    asyncio.run(concurrent_runs())
